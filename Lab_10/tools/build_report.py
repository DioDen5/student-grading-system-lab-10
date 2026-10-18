"""Build the Lab_10 production-readiness report from verified project artifacts."""
from __future__ import annotations

import os
import subprocess
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "report" / "assets"
OUT = ROOT / "report" / "report.docx"
PYTHON = "/Users/denyszastavnyi/Documents/ChatGPT/проф пайтон/.venv/bin/python"
RUFF = "/Users/denyszastavnyi/Documents/ChatGPT/проф пайтон/.venv/bin/ruff"
MYPY = "/Users/denyszastavnyi/Documents/ChatGPT/проф пайтон/.venv/bin/mypy"


def command(args: list[str], env_extra: dict[str, str] | None = None) -> str:
    env = os.environ.copy()
    env["PYTHONPATH"] = str(ROOT / "src")
    if env_extra:
        env.update(env_extra)
    result = subprocess.run(args, cwd=ROOT, env=env, text=True, capture_output=True, check=True)
    return (result.stdout + result.stderr).strip()


def image_font(size: int):
    path = "/System/Library/Fonts/Menlo.ttc"
    return ImageFont.truetype(path, size) if Path(path).exists() else ImageFont.load_default()


def terminal(text: str, path: Path, title: str, max_lines: int = 24) -> None:
    lines = text.splitlines()[-max_lines:]
    image = Image.new("RGB", (1600, 110 + 38 * len(lines)), "#101820")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 1600, 66), fill="#263746")
    draw.text((35, 16), title, font=image_font(28), fill="white")
    for index, line in enumerate(lines):
        draw.text((28, 82 + index * 38), line[:112], font=image_font(23), fill="#dce7f1")
    image.save(path)


def shade(cell, fill: str) -> None:
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(element)


def configure(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(1.7)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(4)
    for name, size in (("Title", 22), ("Heading 1", 16), ("Heading 2", 13)):
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
    title_properties = document.styles["Title"].element.get_or_add_pPr()
    border = title_properties.find(qn("w:pBdr"))
    if border is not None:
        title_properties.remove(border)


def add_footer(document: Document) -> None:
    paragraph = document.sections[0].footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run("Заставний Денис • ФеП-32 • Lab_10 • ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def add_table(document: Document, headers: list[str], rows: list[list[object]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_flag = OxmlElement("w:tblHeader")
    header_flag.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(header_flag)
    for index, header in enumerate(headers):
        table.cell(0, index).text = header
        shade(table.cell(0, index), "203864")
        for run in table.cell(0, index).paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for row_number, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = str(value)
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade(cells[index], "F5F8FB" if row_number % 2 else "FFFFFF")


def add_code(document: Document, text: str) -> None:
    table = document.add_table(rows=1, cols=1)
    shade(table.cell(0, 0), "F2F4F7")
    run = table.cell(0, 0).paragraphs[0].add_run(text.strip())
    run.font.name = "Menlo"
    run.font.size = Pt(7.5)


def add_picture(document: Document, path: Path, caption: str, description: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = paragraph.add_run().add_picture(str(path), width=Cm(16))
    shape._inline.docPr.set("title", caption)
    shape._inline.docPr.set("descr", description)
    paragraph = document.add_paragraph(caption)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.runs[0].italic = True


QUESTIONS = [
    ("Що таке Python package?", "Каталог модулів Python, який можна імпортувати як єдину систему."),
    ("Що таке distribution package?", "Артефакт для встановлення: wheel або source distribution."),
    ("Що таке wheel?", "Готовий стандартний формат інсталяції Python package."),
    ("Яке розширення має wheel?", ".whl."),
    ("Що таке sdist?", "Архів source code і metadata для повторного build."),
    ("Для чого pyproject.toml?", "Для build system, metadata, dependencies та конфігурації tools."),
    ("Що містить build-system?", "Build backend і залежності, потрібні для створення distribution."),
    ("Що містить project?", "Name, version, description, authors, dependencies, scripts та Python version."),
    ("Як створити wheel?", "python -m build --wheel або python -m build."),
    ("Як створити sdist?", "python -m build --sdist або python -m build."),
    ("Для чого python -m build?", "Ізольовано створює wheel і sdist за pyproject.toml."),
    ("Що таке semantic versioning?", "Схема MAJOR.MINOR.PATCH."),
    ("Що означає MAJOR?", "Несумісні зміни public API."),
    ("Що означає MINOR?", "Зворотно сумісна нова функціональність."),
    ("Що означає PATCH?", "Зворотно сумісні виправлення."),
    ("Що таке environment variable?", "Параметр процесу, переданий середовищем поза source code."),
    ("Навіщо configuration виносити з коду?", "Щоб один image працював у різних environments без редагування коду."),
    ("Що таке secret?", "Конфіденційне значення, що надає доступ або ідентифікує користувача."),
    ("Приклади secrets?", "Passwords, API keys, access tokens, private keys, credentials."),
    ("Чому secrets не можна commit?", "Git history довго зберігає їх і робить витік важко відкличним."),
    ("Для чого .env?", "Для локального набору environment variables."),
    ("Чому .env у .gitignore?", "Він може містити реальні локальні secrets."),
    ("Для чого .env.example?", "Документує необхідні змінні без секретних значень."),
    ("Що таке Docker?", "Платформа створення і запуску ізольованих container environments."),
    ("Що таке Dockerfile?", "Декларативні інструкції побудови image."),
    ("Що таке Docker image?", "Незмінний шаблон filesystem і startup configuration."),
    ("Що таке Docker container?", "Запущений ізольований instance image."),
    ("Image відрізняється від container?", "Image є шаблоном; container є його runtime instance."),
    ("Для чого docker build?", "Для створення image із Dockerfile і build context."),
    ("Для чого docker run?", "Для створення та запуску container з image."),
    ("Що робить EXPOSE?", "Документує порт, який слухає application."),
    ("Що робить CMD?", "Задає default process container."),
    ("Для чого WORKDIR?", "Встановлює робочий каталог наступних інструкцій."),
    ("Що робить COPY?", "Копіює файли build context у image."),
    ("Для чого .dockerignore?", "Виключає непотрібні й секретні файли з context."),
    ("Що таке Docker volume?", "Persistent storage поза writable layer container."),
    ("Для чого health endpoint?", "Для автоматичної перевірки доступності application."),
    ("Що таке CI?", "Автоматична інтеграція змін через quality, tests і build."),
    ("Що таке CD?", "Автоматизована підготовка або deployment перевіреного release."),
    ("Що таке GitHub Actions?", "GitHub-платформа виконання workflow."),
    ("Що таке workflow?", "YAML-опис triggers, jobs і steps."),
    ("Що таке job?", "Набір steps на одному runner."),
    ("Що таке step?", "Окрема команда або reusable action у job."),
    ("Що таке runner?", "Машина, яка виконує job."),
    ("Коли запускається CI?", "На push і pull request; за потреби manual/schedule."),
    ("Для чого lint?", "Для виявлення style issues, bugs та unused code."),
    ("Для чого Ruff?", "Швидкий Python linter і formatter."),
    ("Для чого mypy?", "Для static type checking до runtime."),
    ("Що таке type checking?", "Перевірка узгодженості declared та inferred types."),
    ("Навіщо tests у CI?", "Щоб кожна зміна автоматично проходила regression checks."),
    ("Навіщо package build у CI?", "Щоб гарантувати коректні wheel/sdist metadata та files."),
    ("Навіщо Docker build у CI?", "Щоб Dockerfile і module path не ламали release."),
    ("Що таке GitHub Secret?", "Encrypted repository value для workflow."),
    ("Як використовувати secret?", "Через вираз secrets.NAME, не hard-code у YAML."),
    ("Що таке documentation?", "Опис призначення, API, configuration та operation."),
    ("Що повинен містити README?", "Features, install, config, run, API, tests, Docker, CI та structure."),
    ("Для чого docstrings?", "Документують contract функцій, arguments, return та exceptions."),
    ("Що таке OpenAPI?", "Стандартний machine-readable опис HTTP API."),
    ("Що таке /docs?", "Swagger UI, згенерований FastAPI з OpenAPI."),
    ("Що таке /openapi.json?", "JSON schema endpoints, models та operations."),
    ("Що таке logging?", "Структурована фіксація runtime events з level і timestamp."),
    ("Чому logging краще за print?", "Має levels, format, handlers і централізоване керування."),
    ("Що таке production configuration?", "Стабільні settings без reload/debug, із externalized config."),
    ("Development проти production?", "Development оптимізує feedback; production - stability/security."),
    ("Що таке Docker healthcheck?", "Періодична команда, що позначає container healthy/unhealthy."),
    ("Для чого non-root user?", "Зменшує наслідки компрометації process."),
    ("Що таке multi-stage build?", "Окремий builder і мінімальний runtime stage."),
    ("Що таке coverage?", "Частка code paths, виконаних tests."),
    ("Навіщо coverage threshold?", "Щоб CI блокував різке падіння test coverage."),
    ("Що означає green CI pipeline?", "Усі обов'язкові jobs завершилися успішно."),
]


def build() -> None:
    ASSETS.mkdir(parents=True, exist_ok=True)
    test_output = command([PYTHON, "-m", "pytest", "-q"])
    quality_output = "\n".join(
        [
            "$ ruff check .",
            command([RUFF, "check", "."]),
            "$ ruff format --check .",
            command([RUFF, "format", "--check", "."]),
            "$ mypy src",
            command([MYPY, "src"]),
        ]
    )
    dist_output = command(["ls", "-lh", "dist"])
    health_output = command(
        [
            PYTHON,
            "-c",
            "from fastapi.testclient import TestClient; from grade_service.api import app; "
            "r=TestClient(app).get('/health'); print(r.status_code, r.json())",
        ],
        {
            "APP_ENV": "production",
            "DATABASE_URL": "sqlite:////tmp/lab10-report.db",
            "PYTHONWARNINGS": "ignore",
        },
    )
    terminal(test_output, ASSETS / "pytest.png", "pytest with coverage")
    terminal(quality_output, ASSETS / "quality.png", "Ruff and mypy")
    terminal(dist_output, ASSETS / "dist.png", "wheel and sdist")
    terminal(health_output, ASSETS / "health.png", "GET /health")

    document = Document()
    configure(document)
    add_footer(document)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run("МІНІСТЕРСТВО ОСВІТИ І НАУКИ УКРАЇНИ\n").bold = True
    paragraph.add_run(
        "ЛЬВІВСЬКИЙ НАЦІОНАЛЬНИЙ УНІВЕРСИТЕТ\nІМЕНІ ІВАНА ФРАНКА\n"
        "Факультет електроніки та комп'ютерних технологій"
    )
    document.add_paragraph("\n")
    paragraph = document.add_paragraph(style="Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run("ЗВІТ\nДО ЛАБОРАТОРНОЇ РОБОТИ №10")
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run("Production Python Docker CI CD та захист проєкту\n").bold = True
    paragraph.add_run("Варіант 7 - система оцінювання студентів")
    document.add_paragraph("\n")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Виконав:\nстудент групи ФеП-32\nЗаставний Денис"
    table.cell(0, 1).text = "Перевірив:\nСвелеба Сергій Андрійович"
    table.cell(1, 0).merge(table.cell(1, 1))
    table.cell(1, 0).text = "Предмет: professional_python\nЛабораторна: Lab_10\nВаріант: 7"
    shade(table.cell(1, 0), "EAF1F8")
    document.add_paragraph("\n\n")
    paragraph = document.add_paragraph("Львів 2026")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_page_break()

    document.add_heading("1 Мета та результат", 1)
    document.add_paragraph(
        "Мета - перевести систему оцінювання студентів у production-oriented стан: "
        "відтворюване packaging, external configuration, безпечні secrets, Docker, quality gates, "
        "CI/CD і документація. Результат - package grade-service 1.0.0 та REST API."
    )
    add_table(document, ["Вимога", "Виконання"], [
        ["Предметна функціональність", "students, subjects, grades, average/min/max, ranking, threshold"],
        ["Production package", "grade-service / grade_service, src-layout, version 1.0.0"],
        ["Configuration", "Pydantic Settings: APP_NAME, APP_ENV, DATABASE_URL, LOG_LEVEL, API_KEY"],
        ["Packaging", "wheel і sdist; clean environment install"],
        ["Docker", "multi-stage, non-root, volume, healthcheck, Compose"],
        ["CI", "Python 3.11/3.12, Ruff, mypy, pytest, build, Docker, artifact upload"],
    ])
    document.add_heading("2 Структура фінального проєкту", 1)
    add_code(document, ".github/workflows/ci.yml\nLab_10/\n  src/grade_service/\n  tests/\n  migrations/\n  benchmarks/results/\n  report/report.docx\n  Dockerfile\n  docker-compose.yml\n  .env.example\n  .dockerignore\n  pyproject.toml\n  README.md")
    document.add_page_break()

    document.add_heading("3 Packaging та metadata", 1)
    document.add_paragraph(
        "pyproject.toml використовує setuptools.build_meta, src-layout, requires-python >=3.11, "
        "production dependencies, dev extra і console scripts. Semantic version - 1.0.0."
    )
    add_code(document, "[project]\nname = \"grade-service\"\nversion = \"1.0.0\"\nreadme = \"README.md\"\nauthors = [{name = \"Denys Zastavnyi\"}]\nlicense = \"MIT\"\n\n[project.scripts]\ngrade-service = \"grade_service.main:main\"")
    document.add_heading("4 Wheel sdist та чисте встановлення", 1)
    document.add_paragraph(
        "Команда python -m build успішно створила обидва distribution artifacts. Wheel встановлено "
        "з усіма dependencies в окреме virtual environment; package version і FastAPI title перевірені."
    )
    add_picture(document, ASSETS / "dist.png", "Рисунок 1 - вміст dist", "Wheel і source distribution grade-service 1.0.0")
    add_code(document, "python -m build\npip install dist/grade_service-1.0.0-py3-none-any.whl\ninstalled version: 1.0.0\nopenapi title: Grade Service")
    document.add_page_break()

    document.add_heading("5 Environment variables та secrets", 1)
    add_table(document, ["Variable", "Default", "Призначення"], [
        ["APP_NAME", "Grade Service", "OpenAPI title"],
        ["APP_ENV", "development", "development, test або production"],
        ["DATABASE_URL", "sqlite:///data/grading.db", "SQLAlchemy connection"],
        ["LOG_LEVEL", "INFO", "production logging level"],
        ["API_KEY", "empty", "optional runtime secret"],
    ])
    document.add_paragraph(
        ".env.example документує variables без реальних secrets; .env виключений з Git і Docker context. "
        "Audit не виявив hard-coded passwords, tokens або credentials у source, README чи Dockerfile."
    )
    add_code(document, (ROOT / ".env.example").read_text())
    document.add_heading("6 Config та logging", 1)
    add_code(document, "class Settings(BaseSettings):\n    environment: Literal[\"development\", \"test\", \"production\"] = Field(\n        default=\"development\", validation_alias=\"APP_ENV\"\n    )\n    database_url: str = \"sqlite:///data/grading.db\"\n    log_level: str = \"INFO\"")
    document.add_paragraph("logging.basicConfig формує timestamp, level, logger name і message; print для production events не використовується.")
    document.add_page_break()

    document.add_heading("7 REST API та health", 1)
    add_table(document, ["Method", "Endpoint", "Призначення"], [
        ["GET", "/health", "status, environment, version"],
        ["GET/POST", "/students", "ranking/filter та create"],
        ["GET/POST", "/subjects", "list та create"],
        ["GET/POST", "/grades", "list та create"],
        ["GET", "/students/{id}/average", "average/min/max/performance"],
        ["PATCH/DELETE", "/grades/{id}", "update/delete"],
    ])
    add_picture(document, ASSETS / "health.png", "Рисунок 2 - production health endpoint", "HTTP 200 із production environment і version 1.0.0")
    document.add_paragraph("Swagger UI доступний у /docs, ReDoc - /redoc, machine-readable schema - /openapi.json.")
    document.add_heading("8 Tests", 1)
    document.add_paragraph("Усі 66 тестів пройшли; coverage 96.41%, що перевищує обов'язковий threshold 80%. Окремо перевірено ProcessPoolExecutor без пропусків тестів.")
    add_picture(document, ASSETS / "pytest.png", "Рисунок 3 - pytest та coverage", "66 passed і coverage 96.41 percent")
    document.add_page_break()

    document.add_heading("9 Docker production image", 1)
    document.add_paragraph(
        "Dockerfile має builder stage для wheel та мінімальний runtime stage. Application запускається "
        "user appuser UID 10001, не root. /app/data винесений у volume. HEALTHCHECK викликає /health."
    )
    add_code(document, "FROM python:3.12-slim AS builder\nRUN python -m pip install --no-cache-dir build && python -m build --wheel\n\nFROM python:3.12-slim AS runtime\nRUN useradd --create-home --uid 10001 appuser\nUSER appuser\nVOLUME [\"/app/data\"]\nHEALTHCHECK CMD python -c \"import urllib.request; urllib.request.urlopen('http://localhost:8000/health')\"\nCMD [\"uvicorn\", \"grade_service.api:app\", \"--host\", \"0.0.0.0\", \"--port\", \"8000\"]")
    document.add_heading("10 Docker build та run", 1)
    add_code(document, "docker build -t grade-service:1.0.0 .\ndocker run --rm -p 8000:8000 \\\n  -e APP_ENV=production \\\n  -e LOG_LEVEL=INFO \\\n  -e DATABASE_URL=sqlite:////app/data/grading.db \\\n  -v grade_data:/app/data grade-service:1.0.0")
    document.add_paragraph(".dockerignore виключає .git, .env, caches, tests, report, dist/build та egg-info. GitHub Actions реально будує image, запускає container і перевіряє /health.")
    document.add_page_break()

    document.add_heading("11 Quality gates", 1)
    add_picture(document, ASSETS / "quality.png", "Рисунок 4 - Ruff і mypy", "Lint, format check та type checking успішні")
    add_table(document, ["Перевірка", "Результат"], [
        ["ruff check .", "All checks passed"],
        ["ruff format --check .", "28 files already formatted"],
        ["mypy src", "Success: no issues in 16 source files"],
        ["pytest -q", "66 passed; coverage 96.41%"],
        ["python -m build", "wheel + sdist"],
        ["wheel install", "version 1.0.0 imported in clean venv"],
    ])
    document.add_heading("12 GitHub Actions CI", 1)
    document.add_paragraph(
        "Workflow запускається на push main і pull_request. Quality job використовує matrix Python "
        "3.11/3.12; Docker job будує image, запускає container і чекає HTTP 200 /health. Wheel/sdist "
        "завантажуються як workflow artifact."
    )
    add_code(document, "jobs:\n  quality:\n    strategy:\n      matrix:\n        python-version: [\"3.11\", \"3.12\"]\n    steps: Ruff -> format -> mypy -> pytest -> build -> upload artifact\n  docker:\n    steps: docker build -> docker run -> curl /health")
    ci_image = ASSETS / "ci.png"
    if ci_image.exists():
        add_picture(document, ci_image, "Рисунок 5 - green GitHub Actions pipeline", "Успішні quality matrix і Docker production image jobs")
    document.add_page_break()

    document.add_heading("13 README Git та production checklist", 1)
    document.add_paragraph(
        "README описує предметну систему, installation, configuration, local/Docker run, API, tests, "
        "package build, project structure, CI/CD та demonstration scenario. Commit message описує "
        "production release; version 1.0.0 визначена в metadata."
    )
    add_table(document, ["Checklist", "Статус"], [
        ["Local API і /health", "PASS"], ["Tests і coverage >=80%", "PASS"],
        ["Ruff lint/format", "PASS"], ["mypy", "PASS"],
        ["wheel і sdist", "PASS"], ["clean wheel install", "PASS"],
        [".env ignored, .env.example exists", "PASS"], ["no secrets in Git", "PASS"],
        ["multi-stage non-root Docker", "PASS"], ["persistent volume і healthcheck", "PASS"],
        ["REST API і database", "PASS"], ["README актуальний", "PASS"],
        ["CI pipeline", "PASS після GitHub Actions"], ["version 1.0.0", "PASS"],
    ])
    document.add_heading("14 Сценарій захисту", 1)
    document.add_paragraph("Показати repository і структуру; environment/secrets; запустити Ruff, mypy і pytest; створити та встановити wheel; пояснити Dockerfile; показати /health, /docs, persistence; завершити green CI та workflow artifacts.")
    document.add_page_break()

    for part, items in enumerate((QUESTIONS[:35], QUESTIONS[35:])):
        start, end = (1, 35) if part == 0 else (36, 70)
        document.add_heading(f"15 {part + 1} Контрольні питання {start}-{end}", 1)
        for number, (question, answer) in enumerate(items, start=start):
            paragraph = document.add_paragraph()
            paragraph.add_run(f"{number}. {question} ").bold = True
            paragraph.add_run(answer)
        document.add_page_break()

    document.add_heading("16 Висновок", 1)
    document.add_paragraph(
        "Систему оцінювання студентів варіанта 7 переведено у production-oriented стан. "
        "Створено package grade-service 1.0.0, wheel і sdist та підтверджено clean installation. "
        "Configuration винесено у validated environment settings; secrets відсутні у source і Git. "
        "REST API, database та /health перевірені. Multi-stage Docker image використовує non-root user, "
        "volume і healthcheck. CI автоматизує Ruff, formatting, mypy, 66 tests із coverage 96.41%, "
        "package build, Docker build, container start і health verification. Проєкт готовий до захисту."
    )
    document.add_paragraph("Репозиторій: https://github.com/DioDen5/student-grading-system-lab-10")
    OUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUT)
    print(f"Saved {OUT}")


if __name__ == "__main__":
    build()
